#!/usr/bin/env python3
"""Generate readable Word handbook; normative machine contracts stay in YAML files."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import mistune
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'SettleGraph_Backend_Implementation_and_API_Handoff.docx'
NAVY=RGBColor(20,46,73)
BLUE=RGBColor(32,91,143)
TEAL=RGBColor(19,113,111)
GRAY=RGBColor(85,101,116)

D=Document();sec=D.sections[0]
sec.page_height=Inches(11.69);sec.page_width=Inches(8.27)
sec.top_margin=Inches(.71);sec.bottom_margin=Inches(.69)
sec.left_margin=Inches(.80);sec.right_margin=Inches(.76)
sec.header_distance=Inches(.31);sec.footer_distance=Inches(.33)

st=D.styles
normal=st['Normal'];normal.font.name='Aptos';normal.font.size=Pt(9.25);normal.font.color.rgb=NAVY
normal.paragraph_format.space_after=Pt(5);normal.paragraph_format.line_spacing=1.12
for level,size in [(1,15),(2,12),(3,10.5),(4,9.5)]:
 ss=st[f'Heading {level}'];ss.font.name='Aptos Display';ss.font.size=Pt(size);ss.font.bold=True;ss.font.color.rgb=BLUE if level==1 else NAVY
 ss.paragraph_format.space_before=Pt(13 if level==1 else 9);ss.paragraph_format.space_after=Pt(5)
 ss.paragraph_format.keep_with_next=True

# custom Code style
if 'SG Code' not in st:
 ss=st.add_style('SG Code',1)
else:ss=st['SG Code']
ss.font.name='DejaVu Sans Mono';ss.font.size=Pt(7.25);ss.font.color.rgb=RGBColor(47,62,77)
ss.paragraph_format.left_indent=Inches(.13)
ss.paragraph_format.right_indent=Inches(.08)
ss.paragraph_format.space_after=Pt(1)
ss.paragraph_format.line_spacing=1.03

header=sec.header.paragraphs[0]
header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=header.add_run('SETTLEGRAPH  /  ENGINEERING SPECIFICATION')
r.bold=True;r.font.size=Pt(8);r.font.color.rgb=GRAY
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('SettleGraph AI   •   Proposed engineering blueprint     |     ');r.font.size=Pt(7.5);r.font.color.rgb=GRAY
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)

# background shades
def shade(cell_or_para, color='EDF3F8'):
 target=cell_or_para._tc.get_or_add_tcPr() if hasattr(cell_or_para,'_tc') else cell_or_para._p.get_or_add_pPr()
 sh=OxmlElement('w:shd');sh.set(qn('w:fill'),color);target.append(sh)

def keep_row(row):
 trpr=row._tr.get_or_add_trPr();elem=OxmlElement('w:cantSplit');trpr.append(elem)

def set_cell_margin(cell,top=85,start=105,bottom=85,end=105):
 tc=cell._tc;tcPr=tc.get_or_add_tcPr(); margins=OxmlElement('w:tcMar')
 for name,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
  item=OxmlElement('w:'+name);item.set(qn('w:w'),str(val));item.set(qn('w:type'),'dxa');margins.append(item)
 tcPr.append(margins)

def inner_text(children):
 return ''.join((t.get('raw') or inner_text(t.get('children',[]))) for t in children)

def add_inline(p,nodes,bold=False,italic=False):
 for n in nodes:
  t=n['type']
  if t in ('text','codespan','inline_html'):
   r=p.add_run(n.get('raw',''));r.bold=bold;r.italic=italic
   if t=='codespan':r.font.name='DejaVu Sans Mono';r.font.size=Pt(8);r.font.color.rgb=TEAL
  elif t in ('strong','emphasis'):
   add_inline(p,n.get('children',[]),bold or t=='strong',italic or t=='emphasis')
  elif t=='link':
   add_inline(p,n.get('children',[]),bold,italic)
   r=p.add_run(f' ({n.get("attrs",{}).get("url","")})');r.font.size=Pt(7);r.font.color.rgb=GRAY
  elif t=='softbreak':p.add_run(' ')
  elif t=='linebreak':p.add_run('\n')
  elif 'children' in n:add_inline(p,n['children'],bold,italic)

def render_tokens(nodes, depth=0):
 for n in nodes:
  t=n['type']
  if t=='blank_line':continue
  if t=='heading':
   level=n.get('attrs',{}).get('level',2)
   if level>4:level=4
   p=D.add_paragraph(style=f'Heading {level}');add_inline(p,n.get('children',[]))
  elif t=='paragraph':
   p=D.add_paragraph();add_inline(p,n.get('children',[]))
  elif t=='block_code':
   raw=n.get('raw','').rstrip('\n'); lines=raw.splitlines()
   # Long code blocks break across pages naturally rather than truncating in a single cell.
   for idx,line in enumerate(lines):
    # Prefer short visual lines; wrapping retained in Word.
    p=D.add_paragraph(style='SG Code');shade(p,'F1F5F9');p.add_run(line)
    if idx==0:p.paragraph_format.space_before=Pt(4)
   D.add_paragraph().paragraph_format.space_after=Pt(2)
  elif t=='list':
   ordered=n.get('attrs',{}).get('ordered',False)
   for i,item in enumerate(n.get('children',[])):
    par=D.add_paragraph()
    par.paragraph_format.left_indent=Inches(.19+.18*depth)
    par.paragraph_format.first_line_indent=Inches(-.15)
    par.paragraph_format.space_after=Pt(2.2)
    if item['type']=='task_list_item':mark='☑' if item.get('attrs',{}).get('checked') else '☐'
    else:mark=f'{i+1}.' if ordered else '•'
    prefix=par.add_run(mark+'  ');prefix.bold=True;prefix.font.color.rgb=TEAL
    for ch in item.get('children',[]):
     if ch['type'] in ('block_text','paragraph'): add_inline(par,ch.get('children',[]))
     elif ch['type']=='list':render_tokens([ch],depth+1)
  elif t=='table':
   heads=[];rows=[]
   for part in n.get('children',[]):
    if part['type']=='table_head':heads=[inner_text(c.get('children',[])) for c in part.get('children',[])]
    if part['type']=='table_body':
     for row in part.get('children',[]): rows.append([inner_text(c.get('children',[])) for c in row.get('children',[])])
   ncols=max(len(heads),max((len(x) for x in rows),default=0))
   if ncols==0:continue
   # Long URL/code columns may wrap, keep font smaller.
   T=D.add_table(rows=1 if heads else 0,cols=ncols);T.style='Table Grid';T.autofit=True
   if heads:
    for idx,s in enumerate(heads):
     cell=T.rows[0].cells[idx];cell.text=s;shade(cell,'DDEAF3');set_cell_margin(cell)
     for para in cell.paragraphs:
      for run in para.runs:run.bold=True;run.font.color.rgb=NAVY;run.font.size=Pt(7.8)
    keep_row(T.rows[0])
   for ridx,rowdata in enumerate(rows):
    row=T.add_row();keep_row(row)
    for idx in range(ncols):
     cell=row.cells[idx];cell.text=rowdata[idx] if idx<len(rowdata) else ''
     set_cell_margin(cell)
     if ridx%2==1:shade(cell,'F5F8FB')
     for para in cell.paragraphs:
      para.paragraph_format.space_after=Pt(0)
      for run in para.runs:run.font.size=Pt(7.65)
   D.add_paragraph().paragraph_format.space_after=Pt(1)
  elif t=='block_quote':
   p=D.add_paragraph();p.paragraph_format.left_indent=Inches(.22);shade(p,'EDF3F8')
   p.add_run(inner_text(n.get('children',[]))).italic=True
  elif t=='thematic_break':D.add_paragraph('— '*24)
  elif 'children' in n:render_tokens(n['children'],depth)

def add_md(filename,title,newpage=True):
 if newpage:D.add_page_break()
 D.add_heading(title,0)
 content=(ROOT/filename).read_text()
 nodes=mistune.create_markdown(renderer='ast',plugins=['table','task_lists','strikethrough'])(content)
 # Document-specific first title duplicates section heading; keep as children by default for fidelity.
 render_tokens(nodes)

# Cover page
p=D.add_paragraph();p.paragraph_format.space_before=Pt(83)
r=p.add_run('SETTLEGRAPH AI');r.font.name='Aptos Display';r.bold=True;r.font.size=Pt(34);r.font.color.rgb=NAVY
p=D.add_paragraph('BACKEND EXECUTION, PROJECT MANAGEMENT & API CONTRACTS');p.style='Subtitle'
for run in p.runs:run.font.size=Pt(18);run.font.bold=True;run.font.color.rgb=TEAL
p=D.add_paragraph('Complete implementation phases • Financial invariants • BankSim • AWS • Web & Flutter contracts')
for r in p.runs:r.font.size=Pt(12);r.font.color.rgb=GRAY
D.add_paragraph('')
T=D.add_table(rows=4,cols=2);T.style='Table Grid'
cover=[('Edition','2.1 / 24 October 2026'),('Implementation scope','One Go backend engineer; Next.js and Flutter clients'),('API source of truth','contracts/settlegraph-public.openapi.yaml'),('Status','Proposed specification and static fixtures — backend NOT yet built')]
for row,pair in zip(T.rows,cover):
 for cell,val in zip(row.cells,pair):cell.text=val;set_cell_margin(cell,100,150,100,150)
 shade(row.cells[0],'DDEAF3')
 for r in row.cells[0].paragraphs[0].runs:r.bold=True;r.font.color.rgb=NAVY
D.add_paragraph('')
p=D.add_paragraph('Engineering principle');p.style='Heading 2'
p=D.add_paragraph('Simulated banks produce verifiable test observations. A deterministic Go reconciliation core checks payment relationships. AI investigates unresolved discrepancies, but only source evidence and authorized reviewers can support financial decisions.')
p.paragraph_format.space_after=Pt(12)
p=D.add_paragraph('CONTENTS');p.style='Heading 2'
for txt in ['Architecture and trust boundaries','PM charter, capacity, gate approvals, risk register and 22 action tickets','Phases 0–9, step-by-step with acceptance checkboxes','Financial invariants and test matrix','Web/Flutter API handoff and example payloads','All 49 planned HTTP operations','Copy-ready coding-agent execution instructions']:
 p=D.add_paragraph('•  '+txt);p.paragraph_format.space_after=Pt(2)
D.add_paragraph('')
p=D.add_paragraph('Companion files: full machine-readable OpenAPI YAML, JSON schemas, SQL migrations, deterministic fixtures and validation scripts are supplied in the ZIP.');p.runs[0].italic=True

D.add_page_break();D.add_heading('ARCHITECTURE AT A GLANCE',0)
for img,caption in [('system-architecture.png','Figure 1 — Deployment and trust boundary: BankSim has a private database; SettleGraph consumes signed events and authorized provider reads.'),('financial-flow.png','Figure 2 — Distinct financial matching stages; missing evidence becomes a case, not proof of money loss.'),('case-lifecycle.png','Figure 3 — Cases are versioned and human-reviewed; new evidence does not secretly resolve them.')]:
 D.add_heading(caption.split(' — ')[0],2)
 para=D.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
 w=6.45 if 'system' in img else 6.25
 para.add_run().add_picture(str(ROOT/'diagrams'/img),width=Inches(w))
 p=D.add_paragraph(caption);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 for r in p.runs:r.italic=True;r.font.size=Pt(7.5);r.font.color.rgb=GRAY

add_md('PROJECT_MANAGER_DELIVERY_PLAYBOOK.md','PART I — PROJECT DELIVERY & EXECUTION CONTROL')
add_md('BACKEND_IMPLEMENTATION_BLUEPRINT.md','PART II — DETAILED BACKEND IMPLEMENTATION')
add_md('FRONTEND_MOBILE_HANDOFF.md','PART III — FRONTEND & FLUTTER HANDOFF')
add_md('API_REFERENCE.md','PART IV — ALL PLANNED API ROUTES')
add_md('AGENT_IMPLEMENTATION_INSTRUCTIONS.md','PART V — CODING AGENT EXECUTION PROMPT')

# Core props and no accidental metadata disclosure
D.core_properties.title='SettleGraph AI — Project Delivery, Backend Implementation and API Handoff'
D.core_properties.subject='AWS fintech hackathon synthetic-bank reconciliation implementation specification'
D.core_properties.author='SettleGraph Engineering'
D.core_properties.keywords='Go, AWS, Bedrock, BankSim, OpenAPI, reconciliation, Flutter, Next.js'
D.save(OUT)
print('Wrote',OUT,OUT.stat().st_size)
